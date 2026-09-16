"""Read-only Excel inspection. Phase 1: no writes, no formulas.

The first feature this enables is `pepe ask atlas "inspect models/base_lbo.xlsx"`.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl


@dataclass
class WorkbookInfo:
    path: str
    sheets: list[str]
    named_ranges: list[str]
    sheet_dimensions: dict[str, str]

    def to_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "sheets": self.sheets,
            "named_ranges": self.named_ranges,
            "sheet_dimensions": self.sheet_dimensions,
        }


def inspect_workbook(path: Path) -> WorkbookInfo:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        named = []
        for n in wb.defined_names:
            named.append(str(n))
        dims = {name: wb[name].dimensions for name in wb.sheetnames}
        return WorkbookInfo(
            path=str(path),
            sheets=list(wb.sheetnames),
            named_ranges=named,
            sheet_dimensions=dims,
        )
    finally:
        wb.close()


def list_sheets(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def inspect_range(path: Path, sheet: str, range_ref: str) -> list[list[Any]]:
    """Read values in a sheet + range. data_only=True returns last cached values."""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        ws = wb[sheet]
        out: list[list[Any]] = []
        for row in ws[range_ref]:
            out.append([_coerce(c) for c in row])
        return out
    finally:
        wb.close()


def _coerce(cell: Any) -> Any:
    """Convert openpyxl cell values to JSON-friendly primitives."""
    v = cell.value
    if v is None:
        return None
    if isinstance(v, (int, float, bool, str)):
        return v
    return str(v)
